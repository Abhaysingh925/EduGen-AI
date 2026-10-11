import os
from docx import Document

PARAS_PER_PAGE = 15  # Word has no real pages, so we group paragraphs


def process_docx(file) -> dict:
    filename = os.path.basename(getattr(file, "name", file))
    doc = Document(file)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # include table text too
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                paras.append(row_text)
    pages = []
    for i in range(0, len(paras), PARAS_PER_PAGE):
        pages.append({"page": i // PARAS_PER_PAGE + 1,
                      "text": "\n".join(paras[i:i + PARAS_PER_PAGE])})
    return {"filename": filename, "type": "docx", "pages": pages, "empty_pages": []}
