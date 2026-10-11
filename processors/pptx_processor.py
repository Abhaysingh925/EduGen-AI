import os
from pptx import Presentation


def process_pptx(file) -> dict:
    filename = os.path.basename(getattr(file, "name", file))
    prs = Presentation(file)
    pages, empty = [], []
    for num, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(p for p in parts if p.strip()).strip()
        if not text:
            empty.append(num)
            continue
        pages.append({"page": num, "text": text})   # page = slide number
    return {"filename": filename, "type": "pptx", "pages": pages, "empty_pages": empty}
